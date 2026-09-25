# -*- coding: utf-8 -*-
"""Enxuga o deck da Ramasa, corrige os números do caso e cria o slide de
arquitetura com preço peça por peça."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy, re, sys

NAVY  = RGBColor(0x0F,0x0F,0x1E)
GOLD  = RGBColor(0xC9,0xA8,0x4C)
GOLDD = RGBColor(0xA5,0x84,0x2F)
BODY  = RGBColor(0x16,0x18,0x1D)
GREY  = RGBColor(0x9A,0xA1,0xAE)
CARD  = RGBColor(0xF4,0xF5,0xF8)
WHITE = RGBColor(0xFF,0xFF,0xFF)
NEVOA = RGBColor(0xC8,0xCB,0xDB)
LINHA = RGBColor(0xDD,0xE3,0xEC)

ENTRADA = sys.argv[1]; SAIDA = sys.argv[2]
pr = Presentation(ENTRADA)

# ── ferramentas ───────────────────────────────────────────────────────────────
def texto(shape, novo):
    """Troca o texto preservando a formatação do primeiro run."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs: return
    p0.runs[0].text = novo
    for r in list(p0.runs[1:]): r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]): p._p.getparent().remove(p._p)

def acha(slide, trecho):
    for sh in slide.shapes:
        if sh.has_text_frame and trecho in sh.text_frame.text:
            return sh
    return None

def caixa(slide, x, y, w, h, txt, *, tam=11, bold=False, cor=BODY, fonte='Arial',
          espaco=None, alinha=PP_ALIGN.LEFT, entrelinha=1.15):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = alinha
    p.line_spacing = entrelinha
    r = p.add_run(); r.text = txt
    r.font.size = Pt(tam); r.font.bold = bold; r.font.name = fonte
    r.font.color.rgb = cor
    if espaco is not None:
        r.font._rPr.set('spc', str(int(espaco)))
    return tb

def retangulo(slide, x, y, w, h, *, fill=None, linha=None, raio=0.06):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = raio
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if linha is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = linha; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh

# ── 1) slide 14: o caso da Ramasa, com os números de hoje ────────────────────
s = pr.slides[13]
texto(acha(s, '85 dias'), '86 dias')
texto(acha(s, 'da primeira linha de conteúdo'), 'da primeira conversa até hoje — 18 deles entre a primeira linha de conteúdo e o time no app')
texto(acha(s, '32 de 35'), '30 de 37')
texto(acha(s, '776'), '744')
texto(acha(s, 'ações registradas'), 'ações registradas no primeiro mês de uso')
texto(acha(s, '62'), '64')
alvo = acha(s, 'De 03 de julho')
if alvo:
    tf = alvo.text_frame
    if len(tf.paragraphs) > 1 and tf.paragraphs[1].runs:
        tf.paragraphs[1].runs[0].text = 'De 1º de julho a 25 de setembro de 2026'
        for r in list(tf.paragraphs[1].runs[1:]): r._r.getparent().remove(r._r)

# ── 2) preço da implantação: R$ 25.000, com a tabela cheia como âncora ───────
s = pr.slides[18]
texto(acha(s, 'R$ 35.380'), 'R$ 25.000')
texto(acha(s, 'A linha carregada'),
      'A linha carregada, o conteúdo escrito, a marca configurada, a gerência treinada e o time no ar em três semanas. Tabela cheia R$ 51.930 — no pacote, 52% abaixo.')
s = pr.slides[20]
texto(acha(s, 'R$ 35.380'), 'R$ 25.000')
alvo = acha(pr.slides[18], 'Quando precisar: carro novo')
if alvo:
    texto(alvo, 'Quando precisar: carro novo, R$ 2.500 · evolução sob demanda, R$ 350 por hora. Valores para a linha de 5 modelos e 27 acessórios que já está no ar.')

# ── 3) o slide novo: arquitetura com preço peça por peça ─────────────────────
novo = pr.slides.add_slide(pr.slide_layouts[0])
for ph in list(novo.shapes):
    ph._element.getparent().remove(ph._element)

caixa(novo, 0.60, 0.50, 12.10, 0.50, 'O que entra, peça por peça', tam=32, bold=True, cor=NAVY, fonte='Georgia')
caixa(novo, 0.60, 1.41, 12.10, 0.27,
      'Cada parte tem preço de tabela. No pacote de implantação, tudo isso junto sai por R$ 25.000.',
      tam=17, cor=BODY)

TOPO = 1.92; BASE = 6.60

# coluna A — módulo principal, painel escuro
retangulo(novo, 0.60, TOPO, 4.75, BASE - TOPO, fill=NAVY, raio=0.035)
caixa(novo, 0.92, TOPO + 0.28, 4.10, 0.20, 'MÓDULO PRINCIPAL', tam=11, bold=True, cor=GOLD, espaco=300)
caixa(novo, 0.92, TOPO + 0.58, 4.10, 0.28, 'O que já está no ar', tam=17, bold=True, cor=WHITE, fonte='Georgia')

PRINCIPAL = [
    ('Tira-dúvida com IA',                 'R$ 6.000'),
    ('Trilhas de vídeo · 5 carros',        'R$ 7.500'),
    ('Fichas comerciais · 5 carros',       'R$ 5.000'),
    ('One-page e material de estudo',      'R$ 4.200'),
    ('Painel da gerência',                 'R$ 3.300'),
    ('Catálogo de acessórios · 27 itens',  'R$ 2.430'),
    ('Configuração do grupo e das lojas',  'R$ 1.600'),
    ('Documentos da montadora e notícias', 'R$ 1.400'),
    ('PDFs de apoio, vídeos e relatórios', 'R$ 7.850'),
]
y = TOPO + 1.05
for nome, preco in PRINCIPAL:
    caixa(novo, 0.92, y, 3.05, 0.22, nome, tam=11, cor=NEVOA)
    caixa(novo, 3.95, y, 1.10, 0.22, preco, tam=11, bold=True, cor=WHITE, alinha=PP_ALIGN.RIGHT)
    y += 0.355
lin = novo.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.92), Inches(6.08), Inches(4.13), Pt(0.75))
lin.fill.solid(); lin.fill.fore_color.rgb = RGBColor(0x3A,0x3E,0x52); lin.line.fill.background(); lin.shadow.inherit = False
caixa(novo, 0.92, 6.22, 3.05, 0.22, 'Subtotal', tam=11, bold=True, cor=GOLD)
caixa(novo, 3.95, 6.22, 1.10, 0.22, 'R$ 39.280', tam=11, bold=True, cor=GOLD, alinha=PP_ALIGN.RIGHT)

# coluna B — módulos extras
X = 5.65; L = 3.30
caixa(novo, X, TOPO + 0.02, L, 0.20, 'MÓDULOS EXTRAS', tam=11, bold=True, cor=GOLDD, espaco=300)
EXTRAS = [
    ('Condições comerciais',        'R$ 2.400', False),
    ('Jornada do cliente',          'R$ 4.000', False),
    ('Cultura do grupo',            'R$ 2.400', False),
    ('Rituais do mês, por cargo',   'R$ 2.200', False),
    ('Arte da condição do mês','R$ 1.200', False),
    ('Campanha de incentivo',       'R$ 450',   False),
    ('Concorrentes',                'a definir', True),
]
y = TOPO + 0.36
for nome, preco, aberto in EXTRAS:
    retangulo(novo, X, y, L, 0.46, fill=None if aberto else CARD, linha=LINHA if aberto else None, raio=0.10)
    caixa(novo, X + 0.18, y + 0.13, L - 1.15, 0.22, nome, tam=11, bold=True, cor=NAVY)
    caixa(novo, X + L - 1.05, y + 0.13, 0.87, 0.22, preco, tam=11, bold=True,
          cor=GREY if aberto else GOLDD, alinha=PP_ALIGN.RIGHT)
    y += 0.565
SUBTOTAL_Y = 6.22
caixa(novo, X + 0.18, SUBTOTAL_Y, L - 1.15, 0.22, 'Subtotal', tam=11, bold=True, cor=NAVY)
caixa(novo, X + L - 1.05, SUBTOTAL_Y, 0.87, 0.22, 'R$ 12.650', tam=11, bold=True, cor=GOLDD, alinha=PP_ALIGN.RIGHT)

# coluna C — diferenciais + a âncora
X2 = 9.30; L2 = 3.40
caixa(novo, X2, TOPO + 0.02, L2, 0.20, 'DIFERENCIAIS', tam=11, bold=True, cor=GOLDD, espaco=300)
DIF = [
    ('Time de champions', 'Acompanhamento de quem puxa o time dentro de cada loja.'),
    ('Suporte diário na implantação', 'Nas três semanas de carga, todo dia. Depois, quando precisar.'),
    ('Sócias do varejo automotivo', '+25 anos em treinamento, campanha de incentivo, marketing e performance.'),
    ('Engajamento medido', 'Relatório de uso por pessoa, todo mês. A renovação se decide pelo número.'),
]
y = TOPO + 0.36
for titulo, desc in DIF:
    bolinha = novo.shapes.add_shape(MSO_SHAPE.OVAL, Inches(X2), Inches(y + 0.055), Inches(0.10), Inches(0.10))
    bolinha.fill.solid(); bolinha.fill.fore_color.rgb = GOLD; bolinha.line.fill.background(); bolinha.shadow.inherit = False
    caixa(novo, X2 + 0.24, y, L2 - 0.24, 0.22, titulo, tam=11, bold=True, cor=NAVY)
    caixa(novo, X2 + 0.24, y + 0.26, L2 - 0.24, 0.50, desc, tam=10, cor=BODY)
    y += 0.80

retangulo(novo, X2, BASE - 1.22, L2, 1.22, fill=NAVY, raio=0.09)
caixa(novo, X2 + 0.26, BASE - 1.08, L2 - 0.52, 0.18, 'TUDO ISSO, NO PACOTE', tam=10, bold=True, cor=GOLD, espaco=300)
caixa(novo, X2 + 0.26, BASE - 0.84, L2 - 0.52, 0.44, 'R$ 25.000', tam=28, bold=True, cor=WHITE, fonte='Georgia', entrelinha=1.0)
caixa(novo, X2 + 0.26, BASE - 0.32, L2 - 0.52, 0.22,
      'Tabela cheia R$ 51.930 — 52% abaixo.', tam=10, cor=NEVOA)

caixa(novo, 0.60, BASE + 0.14, 12.10, 0.20,
      'Estes preços são de construir e carregar — pagos uma vez. Manter tudo isso no ar é a licença de R$ 900/mês; manter o conteúdo vivo é a curadoria de R$ 2.400/mês. Carro novo: R$ 2.500.',
      tam=10, cor=GREY)
caixa(novo, 10.70, 7.08, 2.03, 0.14, 'Eleva · GSS Academy', tam=9, cor=GREY, alinha=PP_ALIGN.RIGHT)

pr.save(SAIDA)
print('montado ->', SAIDA, '| slides:', len(pr.slides))

# ── 3b) aparos de texto que a leitura visual apontou ────────────────────────
def aparo(indice, trecho, novo):
    sh = acha(pr.slides[indice], trecho)
    if sh: texto(sh, novo)

# espaço duplo no subtítulo do uso diário
for sl in pr.slides:
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for pp in sh.text_frame.paragraphs:
            for r in pp.runs:
                if '  ' in r.text and 'GSS Academy' not in r.text:
                    r.text = re.sub(r'(?<=\S)  +(?=\S)', ' ', r.text)

# concordância
alvo = acha(pr.slides[20], 'Relatórios de uso semanal')
if alvo: texto(alvo, 'Relatório de uso semanal')

# ── 3c) o slide do crescimento: mais loja, outra bandeira, outro grupo ──────
cresce = pr.slides.add_slide(pr.slide_layouts[0])
for ph in list(cresce.shapes):
    ph._element.getparent().remove(ph._element)

caixa(cresce, 0.60, 0.50, 12.10, 0.50, 'Quando o grupo cresce', tam=32, bold=True, cor=NAVY, fonte='Georgia')
caixa(cresce, 0.60, 1.41, 12.10, 0.27,
      'Três situações diferentes, três contas diferentes. A unidade é a bandeira: loja não custa, catálogo custa.',
      tam=17, cor=BODY)

CASOS = [
    ('MAIS UMA LOJA', 'Mesma bandeira, mesmo dono',
     'R$ 400', 'uma vez · R$ 0 por mês',
     'Configurar a unidade, cadastrar o time e separar o painel. A curadoria não muda — é a mesma carta da montadora e a mesma linha de carros. Você paga o que muda, não o que se repete.'),
    ('OUTRA BANDEIRA', 'No mesmo grupo',
     'R$ 8.000', '+ R$ 2.500 por modelo · R$ 1.800 por mês',
     'Aqui não é mais loja: é outro catálogo, com outra carta de montadora para publicar e conferir todo mês. A licença continua uma só, e a fatura sai no CNPJ da bandeira.'),
    ('OUTRO GRUPO', 'Mesma bandeira, outro dono',
     'R$ 14.000', 'R$ 3.300 por mês',
     'O conteúdo de produto já está escrito — descreve o carro da montadora, não a sua operação. O que é seu (condição, time e números) fica isolado no banco, e nenhum outro grupo enxerga.'),
]
LARG = 3.85; GAP = 0.28; TOPO2 = 1.95; ALT = 3.40
for i, (rotulo, sub, valor, complemento, texto_) in enumerate(CASOS):
    x = 0.60 + i * (LARG + GAP)
    retangulo(cresce, x, TOPO2, LARG, 1.02, fill=NAVY, raio=0.07)
    caixa(cresce, x + 0.26, TOPO2 + 0.22, LARG - 0.52, 0.20, rotulo, tam=11, bold=True, cor=GOLD, espaco=300)
    caixa(cresce, x + 0.26, TOPO2 + 0.52, LARG - 0.52, 0.24, sub, tam=12, cor=NEVOA)
    caixa(cresce, x + 0.26, TOPO2 + 1.24, LARG - 0.52, 0.42, valor, tam=27, bold=True, cor=NAVY, fonte='Georgia', entrelinha=1.0)
    caixa(cresce, x + 0.26, TOPO2 + 1.74, LARG - 0.52, 0.24, complemento, tam=11, bold=True, cor=GOLDD)
    caixa(cresce, x + 0.26, TOPO2 + 2.10, LARG - 0.52, 1.10, texto_, tam=11, cor=BODY, entrelinha=1.25)
    lin2 = cresce.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.26), Inches(TOPO2 + 2.02), Inches(LARG - 0.52), Pt(0.75))
    lin2.fill.solid(); lin2.fill.fore_color.rgb = LINHA; lin2.line.fill.background(); lin2.shadow.inherit = False

FAIXA_Y = 5.75
retangulo(cresce, 0.60, FAIXA_Y, 12.13, 1.05, fill=CARD, raio=0.07)
caixa(cresce, 0.92, FAIXA_Y + 0.20, 4.20, 0.22, 'SEM LIMITE DE ACESSOS', tam=11, bold=True, cor=GOLDD, espaco=300)
caixa(creske := cresce, 0.92, FAIXA_Y + 0.50, 11.50, 0.42,
      'Usuários e lojas ilimitados, em todos os planos. Cobrar por pessoa faria a gerência racionar o acesso — e é justamente o uso que prova o valor. A única proteção é automática, aplicada por pessoa e por dia contra uso de robô, e nunca é faturada.',
      tam=11, cor=BODY, entrelinha=1.25)

caixa(cresce, 10.70, 7.08, 2.03, 0.14, 'Eleva · GSS Academy', tam=9, cor=GREY, alinha=PP_ALIGN.RIGHT)

# ── 4) enxugar: saem os slides que a nova substitui ou que repetem ───────────
# 5  A jornada do atendimento      → vira linha da nova (R$ 4.000)
# 6  A pessoa aprende e prova      → vira linha da nova (trilhas + quiz)
# 12 Tudo o que vem · quem vende   → substituídos pela nova
# 13 Tudo o que vem · quem lidera  → substituídos pela nova
# 18 Como trabalhamos juntos       → repete o slide de investimento
# 20 O que a implantação entrega   → substituído pela nova, com preço
FORA = [5, 6, 12, 13, 18, 20]

lista = pr.slides._sldIdLst
ids = list(lista)
for n in sorted(FORA, reverse=True):
    el = ids[n - 1]
    pr.part.drop_rel(el.rId)
    lista.remove(el)

# ── 5) a nova vai para a penúltima posição, antes de "Próximos passos" ───────
# As duas novas ("peça por peça" e "quando o grupo cresce") são as últimas do
# arquivo; entram na frente de "Próximos passos", nessa ordem.
ids = list(lista)
novas = [ids[-2], ids[-1]]
for el in novas:
    lista.remove(el)
# A posição é calculada UMA vez: recalcular len() a cada inserção empurrava a
# segunda para depois de "Próximos passos".
base = len(list(lista)) - 1
for k, el in enumerate(novas):
    lista.insert(base + k, el)

# ── 6) o rodapé volta a contar certo ─────────────────────────────────────────
import re as _re
for n, sl in enumerate(pr.slides, start=1):
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if 'GSS Academy' in t and len(t) < 40:
            # Reescreve o rodapé inteiro. A versão anterior usava um regex com
            # \d* que casava também a string vazia no fim — e colava o número
            # duas vezes ("Eleva · GSS Academy   14   14").
            p0 = sh.text_frame.paragraphs[0]
            if p0.runs:
                p0.runs[0].text = 'Eleva · GSS Academy   ' + str(n)
                for r in list(p0.runs[1:]): r._r.getparent().remove(r._r)
            for pp in list(sh.text_frame.paragraphs[1:]): pp._p.getparent().remove(pp._p)
            break

pr.save(SAIDA)
print('final ->', SAIDA, '| slides:', len(pr.slides))
